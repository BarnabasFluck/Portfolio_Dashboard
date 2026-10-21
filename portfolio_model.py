"""
Calculation engine for the portfolio dashboard.

Everything here mirrors the logic of Portfolio_Analysis_Simple.xlsm:
  * weekly returns from the Data sheet (stocks: price change; Treasury: carry +
    duration/convexity applied to the weekly change in the 5Y yield)
  * annualised mean, volatility, beta (SLOPE vs S&P 500) and CAPM expected return
  * a 1-year Black-Scholes put on the S&P 500 that hedges the equity beta
On top of that it adds what the workbook does not have: risk-profile
allocations, recurring contributions and a multi-year Monte Carlo projection.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import minimize
from scipy.stats import norm

ASSETS = ["LMT", "TXN", "KO", "UST 5Y"]
STOCKS = ["LMT", "TXN", "KO"]


# --------------------------------------------------------------------------- #
# 1. Reading the workbook
# --------------------------------------------------------------------------- #
@dataclass
class WorkbookInputs:
    prices: pd.DataFrame            # Date index; LMT, TXN, KO, UST5Y_yield, SP500
    weights: dict                   # the workbook's own (max-Sharpe) allocation
    rf: float
    erp: float
    weeks_per_year: int
    put_enabled: bool
    put_strike: float
    put_hedge_share: float
    bond: dict = field(default_factory=dict)
    mandate: dict = field(default_factory=dict)


def _num(x, default):
    return float(x) if isinstance(x, (int, float)) else default


def read_workbook(file_bytes: bytes) -> WorkbookInputs:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, keep_vba=False)
    inp, data, val = wb["Inputs"], wb["Data"], wb["Valuation and Bond"]

    rows = []
    for r in data.iter_rows(min_row=6, max_col=6, values_only=True):
        if r[0] is None:
            break
        rows.append(r)
    prices = pd.DataFrame(rows, columns=["Date", "LMT", "TXN", "KO", "UST5Y_yield", "SP500"])
    prices["Date"] = pd.to_datetime(prices["Date"])
    prices = prices.set_index("Date").astype(float)

    weights = {a: _num(inp[f"B{r}"].value, 0.0) for a, r in zip(ASSETS, range(6, 10))}

    return WorkbookInputs(
        prices=prices,
        weights=weights,
        rf=_num(inp["B14"].value, 0.05),
        erp=_num(inp["B15"].value, 0.05),
        weeks_per_year=int(_num(inp["B17"].value, 52)),
        put_enabled=str(inp["B28"].value).strip().lower() == "yes",
        put_strike=_num(inp["B29"].value, 0.8),
        put_hedge_share=_num(inp["B30"].value, 1.0),
        bond=dict(
            face=_num(val["B34"].value, 1000), coupon=_num(val["B35"].value, 0.05),
            ytm=_num(val["B36"].value, 0.05), years=_num(val["B37"].value, 5),
            freq=int(_num(val["B38"].value, 2)), bump_bp=_num(val["B39"].value, 25),
        ),
        mandate=dict(
            target_return=_num(inp["B20"].value, 0.08), max_vol=_num(inp["B21"].value, 0.15),
            max_drawdown=_num(inp["B22"].value, 0.26), min_treasury=_num(inp["B23"].value, 0.2),
            max_weight=_num(inp["B24"].value, 0.4), min_weight=_num(inp["B25"].value, 0.1),
        ),
    )


# --------------------------------------------------------------------------- #
# 2. Market statistics (Data + Portfolio + Valuation and Bond sheets)
# --------------------------------------------------------------------------- #
def bond_price(face, coupon, ytm, years, freq):
    n = int(round(years * freq))
    c, y = face * coupon / freq, ytm / freq
    t = np.arange(1, n + 1)
    return float(np.sum(c / (1 + y) ** t) + face / (1 + y) ** n)


def duration_convexity(bond: dict):
    b = bond
    bump = b["bump_bp"] / 10000
    p0 = bond_price(b["face"], b["coupon"], b["ytm"], b["years"], b["freq"])
    p_up = bond_price(b["face"], b["coupon"], b["ytm"] + bump, b["years"], b["freq"])
    p_dn = bond_price(b["face"], b["coupon"], b["ytm"] - bump, b["years"], b["freq"])
    duration = (p_dn - p_up) / (2 * p0 * bump)
    convexity = (p_up + p_dn - 2 * p0) / (p0 * bump ** 2)
    return duration, convexity


@dataclass
class MarketStats:
    returns: pd.DataFrame       # weekly returns: LMT, TXN, KO, UST 5Y, S&P 500
    exp_ret: pd.Series          # CAPM expected return per asset (annual)
    vol: pd.Series              # annual volatility per asset
    beta: pd.Series
    cov: pd.DataFrame           # annualised covariance of the 4 assets
    sp_vol: float
    duration: float
    convexity: float


def market_stats(wbi: WorkbookInputs) -> MarketStats:
    p, n = wbi.prices, wbi.weeks_per_year
    dur, conv = duration_convexity(wbi.bond)

    rets = pd.DataFrame(index=p.index)
    for s in STOCKS:
        rets[s] = p[s].pct_change()
    y_prev, dy = p["UST5Y_yield"].shift(1) / 100, p["UST5Y_yield"].diff() / 100
    rets["UST 5Y"] = y_prev / n - dur * dy + 0.5 * conv * dy ** 2
    rets["S&P 500"] = p["SP500"].pct_change()
    rets = rets.dropna()

    sp = rets["S&P 500"]
    beta = rets[ASSETS].apply(lambda c: c.cov(sp) / sp.var())
    exp_ret = wbi.rf + beta * wbi.erp
    vol = rets[ASSETS].std() * np.sqrt(n)
    cov = rets[ASSETS].cov() * n
    return MarketStats(rets, exp_ret, vol, beta, cov, float(sp.std() * np.sqrt(n)), dur, conv)


# --------------------------------------------------------------------------- #
# 3. The put on the S&P 500 (Risk sheet, section 2)
# --------------------------------------------------------------------------- #
def put_numbers(strike, sp_vol, rf, erp, T=1.0):
    """Price and expected payoff of a put, both as a fraction of the notional."""
    s = sp_vol * np.sqrt(T)
    d1 = (np.log(1 / strike) + (rf + sp_vol ** 2 / 2) * T) / s
    d2 = d1 - s
    price = strike * np.exp(-rf * T) * norm.cdf(-d2) - norm.cdf(-d1)
    mu = rf + erp
    d1r = (np.log(1 / strike) + (mu + sp_vol ** 2 / 2) * T) / s
    d2r = d1r - s
    exp_payoff = strike * norm.cdf(-d2r) - np.exp(mu * T) * norm.cdf(-d1r)
    return float(price), float(exp_payoff)


# --------------------------------------------------------------------------- #
# 4. Risk profiles
# --------------------------------------------------------------------------- #
@dataclass
class RiskProfile:
    key: str
    label: str
    description: str
    weights: pd.Series
    put_enabled: bool
    put_strike: float
    put_hedge_share: float
    rules: str


def _optimise(stats: MarketStats, bounds, objective, extra_cons=()):
    mu, cov = stats.exp_ret.values, stats.cov.values
    cons = [{"type": "eq", "fun": lambda w: w.sum() - 1}, *extra_cons]
    x0 = np.array([(lo + hi) / 2 for lo, hi in bounds])
    x0 = x0 / x0.sum()
    f = {
        "min_vol": lambda w: w @ cov @ w,
        "max_return": lambda w: -(w @ mu),
    }[objective]
    res = minimize(f, x0, bounds=bounds, constraints=cons, method="SLSQP",
                   options={"ftol": 1e-12, "maxiter": 500})
    w = np.clip(res.x, 0, None)
    return pd.Series(w / w.sum(), index=ASSETS).round(4)


def build_profiles(wbi: WorkbookInputs, stats: MarketStats) -> dict[str, RiskProfile]:
    cov = stats.cov.values

    low = _optimise(stats, bounds=[(0.05, 0.30)] * 3 + [(0.40, 0.60)], objective="min_vol")

    medium = pd.Series(wbi.weights, index=ASSETS)

    max_vol_high = 0.20
    high = _optimise(
        stats, bounds=[(0.10, 0.55)] * 3 + [(0.0, 0.10)], objective="max_return",
        extra_cons=[{"type": "ineq", "fun": lambda w: max_vol_high ** 2 - w @ cov @ w}],
    )

    return {
        "Small": RiskProfile(
            "Small", "Small risk", "Capital protection first: lots of Treasury, a deeper put.",
            low, True, 0.90, 1.0,
            "Lowest possible volatility · Treasury 40–60% · each stock 5–30% · "
            "put on the S&P 500 at 90% strike (caps equity losses beyond −10%).",
        ),
        "Medium": RiskProfile(
            "Medium", "Medium risk", "The workbook's own max-Sharpe allocation and put.",
            medium, wbi.put_enabled, wbi.put_strike, wbi.put_hedge_share,
            f"Weights from the Inputs sheet (Solver, max Sharpe) · Treasury ≥ "
            f"{wbi.mandate['min_treasury']:.0%} · put {'on' if wbi.put_enabled else 'off'} "
            f"at {wbi.put_strike:.0%} strike.",
        ),
        "High": RiskProfile(
            "High", "High risk", "Growth first: mostly stocks, no hedge.",
            high, False, wbi.put_strike, 0.0,
            "Highest expected return with volatility ≤ 20% · Treasury 0–10% · "
            "each stock 10–55% · no put (full exposure to market falls and rallies).",
        ),
    }


@dataclass
class ProfileStats:
    exp_return_gross: float
    put_cost: float
    exp_return: float
    vol: float
    beta: float
    equity_beta: float
    sharpe: float
    max_drawdown: float
    put_premium_pct: float      # premium as % of portfolio value, paid each year
    max_loss_sp_fall: float | None
    weekly_returns: pd.Series


def profile_stats(p: RiskProfile, stats: MarketStats, wbi: WorkbookInputs) -> ProfileStats:
    w = p.weights.values
    gross = float(w @ stats.exp_ret.values)
    vol = float(np.sqrt(w @ stats.cov.values @ w))
    beta = float(w @ stats.beta.values)
    eq_beta = float(p.weights[STOCKS] @ stats.beta[STOCKS])

    put_cost = premium_pct = 0.0
    max_loss = None
    if p.put_enabled and p.put_hedge_share > 0:
        price, exp_payoff = put_numbers(p.put_strike, stats.sp_vol, wbi.rf, wbi.erp)
        hedged = eq_beta * p.put_hedge_share
        put_cost = (price * np.exp(wbi.rf) - exp_payoff) * hedged
        premium_pct = price * hedged
        if p.put_hedge_share >= 1:
            max_loss = premium_pct + hedged * (1 - p.put_strike)

    wk = stats.returns[ASSETS] @ w
    growth = (1 + wk).cumprod()
    mdd = float((growth / growth.cummax() - 1).min())
    net = gross - put_cost
    return ProfileStats(gross, put_cost, net, vol, beta, eq_beta,
                        (gross - wbi.rf) / vol, mdd, premium_pct, max_loss, wk)


def stress_test(p: RiskProfile, ps: ProfileStats, stats: MarketStats, scenarios: dict):
    """One-off shocks, same method as the Risk sheet (beta × S&P move; duration/convexity)."""
    out = {}
    for name, (sp_move, dy_bp) in scenarios.items():
        dy = dy_bp / 10000
        asset_ret = stats.beta[STOCKS] * sp_move
        bond_ret = -stats.duration * dy + 0.5 * stats.convexity * dy ** 2
        r = float(p.weights[STOCKS] @ asset_ret + p.weights["UST 5Y"] * bond_ret)
        if p.put_enabled and p.put_hedge_share > 0:
            hedged = ps.equity_beta * p.put_hedge_share
            r += hedged * max(p.put_strike - (1 + sp_move), 0) - ps.put_premium_pct
        out[name] = r
    return out


# --------------------------------------------------------------------------- #
# 5. Projection with contributions
# --------------------------------------------------------------------------- #
FREQ_MONTHS = {"Monthly": 1, "Quarterly": 3, "Yearly": 12}


def contribution_schedule(initial, recurring, freq, years):
    """Cash put in at the end of each month index 0..12*years (index 0 = today)."""
    months = int(round(years * 12))
    flows = np.zeros(months + 1)
    flows[0] = initial
    if recurring > 0 and freq in FREQ_MONTHS:
        step = FREQ_MONTHS[freq]
        flows[step::step] += recurring
    return flows


def deterministic_path(flows, annual_return):
    m = (1 + annual_return) ** (1 / 12) - 1
    v = np.zeros_like(flows)
    v[0] = flows[0]
    for t in range(1, len(flows)):
        v[t] = v[t - 1] * (1 + m) + flows[t]
    return v


def monte_carlo(flows, p: RiskProfile, ps: ProfileStats, stats: MarketStats, wbi: WorkbookInputs,
                n_sims=5000, seed=42):
    """
    Monthly simulation of the S&P 500 and the portfolio (lognormal, correlated through beta).
    The portfolio drift is its CAPM return before the put; the put is then modelled
    explicitly: each year a premium is paid and, if the S&P 500 is below the strike
    at year end, the payoff is received on the hedged amount.
    """
    rng = np.random.default_rng(seed)
    months = len(flows) - 1
    sig_sp, sig_p = stats.sp_vol, ps.vol
    mu_sp, mu_p = wbi.rf + wbi.erp, ps.exp_return_gross
    rho = float(np.clip(ps.beta * sig_sp / sig_p, -0.999, 0.999))

    z1 = rng.standard_normal((n_sims, months))
    z2 = rho * z1 + np.sqrt(1 - rho ** 2) * rng.standard_normal((n_sims, months))
    lr_sp = (np.log(1 + mu_sp) - sig_sp ** 2 / 2) / 12 + sig_sp / np.sqrt(12) * z1
    lr_p = (np.log(1 + mu_p) - sig_p ** 2 / 2) / 12 + sig_p / np.sqrt(12) * z2
    gr_p = np.exp(lr_p)

    hedge = p.put_enabled and p.put_hedge_share > 0
    hedged = ps.equity_beta * p.put_hedge_share if hedge else 0.0

    values = np.zeros((n_sims, months + 1))
    values[:, 0] = flows[0]
    v = np.full(n_sims, flows[0], dtype=float)
    year_start_v = v.copy()
    sp_year = np.ones(n_sims)
    for t in range(1, months + 1):
        if hedge and (t - 1) % 12 == 0:                 # buy this year's put
            year_start_v = v.copy()
            v = v - ps.put_premium_pct * year_start_v
            sp_year = np.ones(n_sims)
        v = v * gr_p[:, t - 1]
        sp_year *= np.exp(lr_sp[:, t - 1])
        if hedge and t % 12 == 0:                       # put expires
            v = v + hedged * year_start_v * np.maximum(p.put_strike - sp_year, 0)
        v = v + flows[t]
        values[:, t] = v
    return values


def money_weighted_return(flows, final_value):
    """Annual IRR of the contribution stream given the final value."""
    months = len(flows) - 1
    if months == 0 or final_value <= 0:
        return np.nan
    t = np.arange(months + 1)

    def npv(r):
        m = (1 + r) ** (1 / 12) - 1
        return final_value / (1 + m) ** months - np.sum(flows / (1 + m) ** t)

    lo, hi = -0.99, 2.0
    if npv(lo) * npv(hi) > 0:
        return np.nan
    for _ in range(200):
        mid = (lo + hi) / 2
        if npv(lo) * npv(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2
